#!/usr/bin/env python3
"""Sisipkan Lampiran E (laporan konsumsi daya baterai) ke docx laporan proyek.
Konten/desain lama tidak disentuh. Angka dibaca dari hasil.json/per_hari.csv analisis_baterai.py."""
import copy, csv, json, os, re, shutil, sys
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

WS = '/home/ibe/sawal_ws'
SRC = f'{WS}/tempp/Laporan-Proyek-AGV-Forklift-2026-09-23.docx'
HERE = os.path.dirname(os.path.abspath(__file__))
LAP = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'laporan_baterai')
shutil.copy2(SRC, os.path.join(HERE, 'Laporan-sebelum-lampiran-E.docx'))
os.chdir(WS)
d = Document(SRC)
body = d.element.body
exec(open(os.path.join(HERE, 'docx_alat2.py')).read())   # heading/para/tabel/fmt/bullet/toc_entries

H = json.load(open(os.path.join(LAP, 'hasil.json')))
per_hari = list(csv.DictReader(open(os.path.join(LAP, 'per_hari.csv'))))
per_sesi = list(csv.DictReader(open(os.path.join(LAP, 'per_sesi.csv'))))
def tgl(k): return f'{int(k[6:8])}/{int(k[4:6])}'
def sesi_lbl(k): return f'{tgl(k[:8])} {k[9:11]}:{k[11:13]}'
def f1(x, nd=1): return fmt(float(x), nd) if x not in (None, '') else '—'

def gambar(nama, caption):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(LAP, nama), width=Cm(16.0))
    c = d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55); _sz(r, 18)
    c.paragraph_format.space_after = Pt(8)

sag = H['sag'] or {}
b = H['banding']

# ================================================================= TULIS
heading('Lampiran E. Laporan konsumsi daya baterai', 1)
para(f'Lampiran ini menjawab pertanyaan "berapa daya yang dipakai robot" dari data yang sudah terekam data_logger '
     f'selama proyek ({tgl(H["hari_pertama"])}–{tgl(H["hari_terakhir"])}, {H["n_hari"]} hari kerja, '
     f'{H["n_sesi_tegangan"]} sesi bertegangan dari {H["n_sesi_total"]} sesi). Robot TIDAK memiliki sensor arus baterai, '
     f'sehingga seluruh angka daya di sini adalah estimasi dari besaran yang tersedia; batasannya dijelaskan di E.1 dan '
     f'wajib dibaca sebelum angka-angkanya dikutip.')

heading('E.1 Sumber data dan batasan', 2)
tabel(['Besaran', 'Yang sebenarnya terekam', 'Akibat untuk laporan ini'], [
    ['Tegangan baterai (kolom battery_v, topik /battery_voltage)',
     'Tegangan DC-link drive MDSM dari TPDO 0x381 (resolusi 0,1 V), diteruskan firmware di frame 0x42A dan '
     'dipublikasikan agv_can_bridge/can_listener sebagai /battery_voltage. Medan battery_voltage dan battery_current '
     'firmware (0x42A byte 0–3) selalu 0,00 karena tidak ada sensor baterai.',
     'Dipakai sebagai proksi tegangan pak baterai 48 V. Turun sesaat saat beban (jatuh tegangan kabel + drive), '
     'jadi bukan pengukur state-of-charge yang presisi; tren antar hari tetap terbaca.'],
    ['Arus (kolom trq_l/trq_r atau wheel_l_eff/wheel_r_eff)',
     'Arus motor per kanal dari TPDO 0x381 (satuan 0,1 A), dijumlahkan |L|+|R|. Tersedia lewat /motor_torque '
     f'(susunan pemetaan, {H["n_sesi_arus_trq"]} sesi) atau effort /joint_states (susunan navigasi, {H["n_sesi_arus_eff"]} sesi) '
     'sejak dekoder 0x381 diperbaiki 14/8. Sesi sebelum itu arusnya kosong/0.',
     'Ini arus MOTOR (fasa), bukan arus yang ditarik dari baterai. Daya = V_dc × I_motor karenanya adalah '
     'BATAS ATAS kasar daya traksi, dan hanya traksi: pompa hidrolik garpu, PC, LiDAR, kamera, relai TIDAK termasuk.'],
    ['Energi (Wh)', 'Σ V·I·Δt pada 50 Hz, dipisah "gerak" (|v_odom| > 1 cm/s) dan "diam".',
     'Angka energi hanya sah untuk perbandingan relatif antar sesi/hari; bukan Wh yang keluar dari baterai.'],
    ['Pembacaan < 40 V', f'Terjadi di {H.get("n_sesi_dclink_rendah", 0)} sesi: DC-link jatuh sampai 19–38 V saat daya drive diputus (ESTOP/relai/boot), bukan baterai kosong.',
     'Sampel < 40 V dikeluarkan dari statistik tegangan; hanya nilai ≥ 40 V yang dianggap tegangan baterai.'],
    ['Jarak tempuh (odom_dist_m)', f'{H.get("n_jarak_buang", 0)} sesi memuat lompatan odometri (rata-rata > 0,7 m/s atau > 2 km per sesi, mis. bug pose saat boot yang diperbaiki 12/8).',
     'Jarak sesi tersebut dinolkan; Wh/m hanya dihitung dari sesi dengan jarak sah.'],
    ['Kapasitas baterai (Ah)', 'Tidak ada di repositori maupun dokumen.',
     'Sisa jam operasi tidak dapat dihitung; hanya laju penurunan tegangan per jam yang bisa dilaporkan.'],
], [2300, 3800, 3311])

heading('E.2 Angka kunci', 2)
rows = [
    ['Rentang data', f'{tgl(H["hari_pertama"])}–{tgl(H["hari_terakhir"])}; {H["n_hari"]} hari; {H["n_sesi_tegangan"]} sesi bertegangan; {H["n_sesi_arus"]} sesi berarus', 'ringkasan_sesi.csv'],
    ['Jam log / jam bergerak / jam diam (sesi berarus)', f'{f1(H["tot_jam_log"])} jam log; gerak {f1(H["jam_gerak"])} jam; diam {f1(H["jam_diam"])} jam', 'per_sesi.csv'],
    ['Jarak tempuh total tercatat', f'{fmt(H["tot_jarak_m"]/1000, 2)} km (dengan data arus: {fmt(H["jarak_dgn_arus_m"]/1000, 2)} km)', 'odom_dist_m'],
    ['Tegangan tertinggi / terendah yang pernah terekam', f'{f1(H["v_tertinggi"])} V / {f1(H["v_terendah"])} V (sesi {sesi_lbl(H["sesi_v_terendah"])})', 'battery_v'],
    ['Arus traksi median saat bergerak / saat diam', f'{f1(H["i_gerak_median"])} A / {f1(H["i_diam_median"])} A', 'sesi ≥ 5 m'],
    ['Arus traksi puncak', f'{f1(H["i_maks"])} A (sesi {sesi_lbl(H["sesi_i_maks"])})', '0x381'],
    ['Daya traksi median saat bergerak (batas atas)', f'{fmt(H["w_gerak_median"], 0)} W', 'V×I'],
    ['Daya "diam" median (drive menahan posisi)', f'≈ {fmt(H["w_diam_median"], 0)} W ({f1(H["i_diam_median"])} A × 48 V)', 'V×I'],
    ['Energi traksi per meter saat bergerak', f'median {fmt(H["wh_per_m_median"], 2)} Wh/m (p25 {fmt(H["wh_per_m_p25"], 2)}; p75 {fmt(H["wh_per_m_p75"], 2)}) — sesi ≥ 10 m', 'per_sesi.csv'],
    ['Energi traksi kumulatif gerak / diam (sesi berarus)', f'{fmt(H["wh_gerak"], 0)} Wh / {fmt(H["wh_diam"], 0)} Wh', 'Σ V·I·Δt'],
    ['Otonom vs joystick (sesi ≥ 10 m)',
     f'otonom n={b["otonom"]["n"]}: {f1(b["otonom"]["wh_per_m"], 2)} Wh/m, {f1(b["otonom"]["i_gerak"])} A, {f1(b["otonom"]["v_ms"], 2)} m/s; '
     f'joystick n={b["joystick"]["n"]}: {f1(b["joystick"]["wh_per_m"], 2)} Wh/m, {f1(b["joystick"]["i_gerak"])} A, {f1(b["joystick"]["v_ms"], 2)} m/s',
     'per_sesi.csv'],
]
if sag:
    rows.append(['Jatuh tegangan terhadap arus (regresi V = V0 − R·I pada sampel 1 Hz)',
                 f'V0 {fmt(sag["v0"], 2)} V; R {fmt(sag["ohm"]*1000, 1)} mΩ; n {fmt(sag["n"], 0)}; korelasi r {fmt(sag["r"], 3)} → '
                 + ('hubungan tidak terukur pada resolusi 0,1 V' if abs(sag["r"]) < 0.1 else 'hubungan lemah tetapi terlihat'),
                 'deret_tegangan_1hz.csv'])
tabel(['Besaran', 'Nilai', 'Sumber'], rows, [2835, 4576, 2000])
para('Cara membaca: arus "diam" yang tidak nol adalah arus penahan drive dalam mode Profile Velocity saat robot berhenti '
     'dengan drive tetap aktif; ia muncul di hampir semua sesi dan menyumbang energi lebih besar daripada gerak karena '
     'robot jauh lebih lama diam daripada berjalan. Perbandingan otonom vs joystick dipengaruhi kecepatan: otonom '
     f'berjalan lebih lambat ({f1(b["otonom"]["v_ms"], 2)} vs {f1(b["joystick"]["v_ms"], 2)} m/s) dan arus geraknya lebih tinggi '
     f'({f1(b["otonom"]["i_gerak"])} vs {f1(b["joystick"]["i_gerak"])} A; banyak koreksi arah dan pivot), sehingga Wh per meter '
     f'{f1(b["otonom"]["wh_per_m"] / b["joystick"]["wh_per_m"], 1)}× lebih besar.', catatan=True)

heading('E.3 Tegangan dan energi per hari kerja', 2)
tabel(['Tanggal', 'Sesi (berarus)', 'Jam awal→akhir', 'V awal → akhir', 'V min', 'Jarak (m)', 'Menit gerak', 'Wh gerak / diam', 'Wh/m'],
      [[tgl(r['tanggal']), f'{r["sesi"]} ({r["sesi_dgn_arus"]})', f'{r["jam_awal"]}→{r["jam_akhir"]}',
        f'{f1(r["v_awal"])} → {f1(r["v_akhir"])}', f1(r['v_min']), r['jarak_m'], f1(r['menit_gerak']),
        f'{f1(r["wh_gerak"])} / {f1(r["wh_diam"])}', f1(r['wh_per_m'], 2)] for r in per_hari],
      [800, 1000, 1200, 1300, 700, 900, 950, 1500, 1061])
gambar('g1_tegangan_per_hari.png', 'Gambar E.1 Tegangan DC-link pada sesi pertama, sesi terakhir, dan nilai terendah tiap hari kerja.')

heading('E.4 Profil satu hari dan jatuh tegangan terhadap arus', 2)
if os.path.exists(os.path.join(LAP, 'g2_deret_18sep.png')):
    gambar('g2_deret_18sep.png', 'Gambar E.2 Tegangan (biru) dan arus traksi L+R (jingga) sepanjang 18/9; sesi disambung berurutan, garis putus = mulai sesi.')
if os.path.exists(os.path.join(LAP, 'g2b_deret_21agt.png')):
    gambar('g2b_deret_21agt.png', 'Gambar E.3 Profil yang sama untuk 21/8 (hari lapangan terpadat: ESTOP, lokalisasi, halangan, kalibrasi garpu).')
if os.path.exists(os.path.join(LAP, 'g3_v_vs_i.png')):
    gambar('g3_v_vs_i.png', 'Gambar E.4 Tegangan DC-link terhadap arus traksi dari seluruh sampel 1 Hz; garis = regresi linier, titik hitam = median per 2 A.')

heading('E.5 Sesi dengan konsumsi energi terbesar', 2)
rows = []
for r in H['sesi_utama']:
    rows.append([sesi_lbl(r['sesi']), fmt(r['durasi_s'], 0), f1(r['jarak_m']), f'{f1(r["bat_awal_v"])} → {f1(r["bat_akhir_v"])} (min {f1(r["bat_min_v"])})',
                 f'{f1(r["i_rata_gerak_a"])} / {f1(r["i_maks_a"])}', f'{f1(r["energi_gerak_wh"])} / {f1(r["energi_diam_wh"])}',
                 f1(r['energi_gerak_wh'] / r['jarak_m'], 2) if r['jarak_m'] >= 5 else '—',
                 'otonom' if (r['detik_otonom'] or 0) > 0.5 * (r['detik_gerak'] or 1) else 'joystick/pemetaan'])
tabel(['Sesi', 'Durasi (s)', 'Jarak (m)', 'V awal → akhir', 'I gerak / puncak (A)', 'Wh gerak / diam', 'Wh/m', 'Kemudi'],
      rows, [1250, 850, 850, 1900, 1400, 1300, 700, 1161])
gambar('g4_wh_per_m.png', 'Gambar E.5 Energi traksi per meter untuk setiap sesi ≥ 10 m (biru = dominan otonom, abu-abu = joystick/pemetaan).')

heading('E.6 Kesimpulan dan rekomendasi', 2)
bullet(f'Pak baterai 48 V bekerja di jendela {f1(H["v_terendah"])}–{f1(H["v_tertinggi"])} V sepanjang proyek; hari kerja biasa '
       f'membuka di ~50 V dan turun 0–2 V sampai sesi terakhir (Tabel E.3). Ambang GUI 17/9 (kuning 47,0 V, merah 45,5 V) '
       f'berada di bawah nilai tipikal hari kerja, jadi peringatan baru akan muncul pada hari yang benar-benar panjang.')
bullet(f'Traksi menarik median {f1(H["i_gerak_median"])} A saat bergerak (puncak {f1(H["i_maks"])} A) dan {f1(H["i_diam_median"])} A saat diam. '
       f'Karena robot diam {f1(H["jam_diam"])} jam vs bergerak {f1(H["jam_gerak"])} jam, energi "diam" ({fmt(H["wh_diam"], 0)} Wh) '
       f'melampaui energi gerak ({fmt(H["wh_gerak"], 0)} Wh). Mematikan penahan drive (disable/quick-stop) saat robot diam '
       f'lama adalah penghematan terbesar yang tersedia tanpa perangkat keras baru.')
bullet(f'Energi traksi per meter median {fmt(H["wh_per_m_median"], 2)} Wh/m. Dengan kecepatan otonom 159 mm/s, itu setara '
       f'≈ {fmt(H["wh_per_m_median"] * 0.159 * 3600, 0)} Wh per jam berjalan (batas atas, traksi saja).')
bullet('Angka di atas TIDAK bisa diubah menjadi "sisa jam operasi" karena kapasitas baterai (Ah) tidak diketahui dan arus '
       'yang terekam adalah arus motor, bukan arus baterai. Rekomendasi: (1) pasang sensor arus baterai (shunt/Hall) di kabel '
       'positif pak — medan battery_voltage/battery_current di frame 0x42A firmware sudah tersedia dan tinggal diisi, '
       '(2) catat kapasitas nominal dan tegangan penuh/kosong pak dari pelat nama, (3) ukur satu siklus penuh dari penuh sampai '
       'ambang merah sambil merekam, sehingga jam operasi dan pemakaian pompa garpu ikut terukur.')
bullet('Berkas data yang dipakai (ringkasan per sesi, per hari, deret tegangan 1 Hz, grafik, dan skrip) disertakan di '
       'folder laporan/baterai/ pada cabang pelaporan repositori, supaya angka di lampiran ini bisa dihitung ulang.')

# ---------------------------------------------------------------- TOC
toc1 = [p for p in d.paragraphs if p.style.name == 'toc 1']
toc2 = [p for p in d.paragraphs if p.style.name == 'toc 2']
tpl1, tpl2 = toc1[-1]._p, toc2[-1]._p
# entri baru diletakkan setelah entri TOC terakhir (level apa pun) yang ada sebelum penutup field
semua_toc = [p._p for p in d.paragraphs if p.style.name in ('toc 1', 'toc 2')]
anchor_after = semua_toc[-1]
for level, text, anchor in toc_entries:
    new = copy.deepcopy(tpl1 if level == 1 else tpl2)
    for el in new.iter():
        for a in list(el.attrib):
            if a.endswith('}paraId') or a.endswith('}textId'): del el.attrib[a]
    hl = new.find(qn('w:hyperlink')); hl.set(qn('w:anchor'), anchor)
    hl.findall('.//' + qn('w:t'))[0].text = text
    hl.find('.//' + qn('w:instrText')).text = f' PAGEREF {anchor} \\h '
    anchor_after.addnext(new); anchor_after = new

d.save(SRC)
print('OK', SRC, 'paragraf', len(d.paragraphs), 'tabel', len(d.tables), 'gambar', len(d.inline_shapes), 'toc baru', len(toc_entries))
